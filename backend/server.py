from dotenv import load_dotenv
from pathlib import Path
import os

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

from fastapi import FastAPI, APIRouter, HTTPException, Depends, UploadFile, File, Form, Request
from fastapi.responses import Response
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
from typing import List, Optional, Dict, Any
import logging
import uuid
import json
import re
import io
import base64
import tempfile
import bcrypt
import jwt
from datetime import datetime, timezone, timedelta

from emergentintegrations.llm.chat import LlmChat, UserMessage, FileContentWithMimeType, ImageContent
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from htmldocx import HtmlToDocx
from bs4 import BeautifulSoup

# ---------------- DB ----------------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_SECRET = os.environ['JWT_SECRET']
JWT_ALGORITHM = "HS256"
EMERGENT_LLM_KEY = os.environ.get('EMERGENT_LLM_KEY')
GEMINI_MODEL = os.environ.get('GEMINI_MODEL', 'gemini-3.1-pro-preview')
GEMINI_IMAGE_MODEL = os.environ.get('GEMINI_IMAGE_MODEL', 'gemini-3.1-flash-image-preview')
SUPERADMIN_EMAIL = os.environ.get('SUPERADMIN_EMAIL', 'superadmin@rpp.com')
SUPERADMIN_PASSWORD = os.environ.get('SUPERADMIN_PASSWORD', 'admin123')

app = FastAPI()
api_router = APIRouter(prefix="/api")

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


# ---------------- Auth helpers ----------------
def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode("utf-8")[:72], bcrypt.gensalt()).decode("utf-8")


def verify_password(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode("utf-8")[:72], hashed.encode("utf-8"))
    except Exception:
        return False


def create_token(user_id: str, email: str) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "exp": datetime.now(timezone.utc) + timedelta(days=7),
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


async def get_current_user(request: Request) -> dict:
    auth = request.headers.get("Authorization", "")
    if not auth.startswith("Bearer "):
        raise HTTPException(status_code=401, detail="Tidak terautentikasi")
    token = auth[7:]
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"id": payload["sub"]}, {"_id": 0, "password_hash": 0})
        if not user:
            raise HTTPException(status_code=401, detail="User tidak ditemukan")
        return user
    except jwt.ExpiredSignatureError:
        raise HTTPException(status_code=401, detail="Sesi berakhir, silakan login kembali")
    except jwt.InvalidTokenError:
        raise HTTPException(status_code=401, detail="Token tidak valid")


# ---------------- Models ----------------
class RegisterInput(BaseModel):
    name: str
    email: EmailStr
    password: str


class LoginInput(BaseModel):
    email: EmailStr
    password: str


class ProfileInput(BaseModel):
    name: Optional[str] = None
    nip: Optional[str] = None
    jabatan: Optional[str] = None
    namaSekolah: Optional[str] = None
    alamatSekolah: Optional[str] = None
    namaKepalaSekolah: Optional[str] = None
    nipKepalaSekolah: Optional[str] = None
    logoMadrasah: Optional[str] = None


class AdminSettingsInput(BaseModel):
    """Pengaturan madrasah yang diatur super admin untuk tiap admin (mengalir ke guru)."""
    namaSekolah: Optional[str] = None
    alamatSekolah: Optional[str] = None
    namaKepalaSekolah: Optional[str] = None
    nipKepalaSekolah: Optional[str] = None
    logoMadrasah: Optional[str] = None


class CreateUserInput(BaseModel):
    name: str
    email: EmailStr
    password: str


class DocumentInput(BaseModel):
    type: str
    title: str
    meta: Dict[str, Any] = Field(default_factory=dict)
    fields: Dict[str, Any] = Field(default_factory=dict)
    content_html: Optional[str] = None


class GenerateInput(BaseModel):
    type: str
    inputs: Dict[str, Any] = Field(default_factory=dict)


def public_user(u: dict) -> dict:
    u.pop("password_hash", None)
    u.pop("_id", None)
    return u


# ---------------- Auth routes ----------------
@api_router.post("/auth/register")
async def register(data: RegisterInput):
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    user = {
        "id": str(uuid.uuid4()),
        "name": data.name,
        "email": email,
        "password_hash": hash_password(data.password),
        "role": "guru",
        "admin_id": None,
        "created_by": None,
        "nip": "",
        "jabatan": "Guru",
        "namaSekolah": "",
        "alamatSekolah": "",
        "namaKepalaSekolah": "",
        "nipKepalaSekolah": "",
        "logoMadrasah": "",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.users.insert_one(user)
    token = create_token(user["id"], email)
    return {"token": token, "user": public_user(dict(user))}


@api_router.post("/auth/login")
async def login(data: LoginInput):
    email = data.email.lower()
    user = await db.users.find_one({"email": email})
    if not user or not verify_password(data.password, user["password_hash"]):
        raise HTTPException(status_code=401, detail="Email atau kata sandi salah")
    token = create_token(user["id"], email)
    return {"token": token, "user": public_user(dict(user))}


@api_router.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user


@api_router.put("/auth/profile")
async def update_profile(data: ProfileInput, user: dict = Depends(get_current_user)):
    update = {k: v for k, v in data.model_dump().items() if v is not None}
    if update:
        await db.users.update_one({"id": user["id"]}, {"$set": update})
    fresh = await db.users.find_one({"id": user["id"]}, {"_id": 0, "password_hash": 0})
    return fresh


# ---------------- Role helpers & admin management ----------------
async def visible_owner_ids(user: dict) -> Optional[List[str]]:
    """Return list of owner_ids the user may see. None means ALL (superadmin)."""
    role = user.get("role")
    if role == "superadmin":
        return None
    if role == "admin":
        gurus = await db.users.find({"admin_id": user["id"]}, {"id": 1, "_id": 0}).to_list(2000)
        return [user["id"]] + [g["id"] for g in gurus]
    return [user["id"]]


def require_role(*roles):
    async def dep(user: dict = Depends(get_current_user)) -> dict:
        if user.get("role") not in roles:
            raise HTTPException(status_code=403, detail="Akses ditolak")
        return user
    return dep


@api_router.post("/admin/users")
async def create_managed_user(data: CreateUserInput, user: dict = Depends(get_current_user)):
    role = user.get("role")
    if role == "superadmin":
        new_role = "admin"
    elif role == "admin":
        new_role = "guru"
    else:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    email = data.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(status_code=400, detail="Email sudah terdaftar")
    doc = {
        "id": str(uuid.uuid4()),
        "name": data.name,
        "email": email,
        "password_hash": hash_password(data.password),
        "role": new_role,
        "admin_id": user["id"] if new_role == "guru" else None,
        "created_by": user["id"],
        "nip": "",
        "jabatan": "Guru" if new_role == "guru" else "Admin",
        "namaSekolah": user.get("namaSekolah", ""),
        "alamatSekolah": user.get("alamatSekolah", ""),
        "namaKepalaSekolah": user.get("namaKepalaSekolah", "") if new_role == "guru" else "",
        "nipKepalaSekolah": user.get("nipKepalaSekolah", "") if new_role == "guru" else "",
        "logoMadrasah": user.get("logoMadrasah", "") if new_role == "guru" else "",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.users.insert_one(doc)
    return public_user(dict(doc))


@api_router.get("/admin/users")
async def list_managed_users(user: dict = Depends(get_current_user)):
    role = user.get("role")
    if role == "superadmin":
        q = {"role": "admin"}
    elif role == "admin":
        q = {"admin_id": user["id"]}
    else:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    users = await db.users.find(q, {"_id": 0, "password_hash": 0}).sort("created_at", -1).to_list(2000)
    # attach doc counts using a single aggregation query (avoids N+1)
    user_ids = [u["id"] for u in users]
    counts = {}
    if user_ids:
        pipeline = [
            {"$match": {"owner_id": {"$in": user_ids}}},
            {"$group": {"_id": "$owner_id", "count": {"$sum": 1}}},
        ]
        async for r in db.documents.aggregate(pipeline):
            counts[r["_id"]] = r["count"]
    for u in users:
        u["doc_count"] = counts.get(u["id"], 0)
    return users


MADRASAH_FIELDS = ["namaSekolah", "alamatSekolah", "namaKepalaSekolah", "nipKepalaSekolah", "logoMadrasah"]


@api_router.put("/admin/users/{user_id}")
async def update_managed_user(user_id: str, data: AdminSettingsInput, user: dict = Depends(get_current_user)):
    role = user.get("role")
    target = await db.users.find_one({"id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    # authorization: superadmin -> admin, admin -> own guru
    if role == "superadmin" and target.get("role") == "admin":
        pass
    elif role == "admin" and target.get("admin_id") == user["id"]:
        pass
    else:
        raise HTTPException(status_code=403, detail="Akses ditolak")

    update = {k: v for k, v in data.model_dump().items() if v is not None}
    if update:
        await db.users.update_one({"id": user_id}, {"$set": update})
        # cascade madrasah settings to all gurus under this admin
        if target.get("role") == "admin":
            cascade = {k: v for k, v in update.items() if k in MADRASAH_FIELDS}
            if cascade:
                await db.users.update_many({"admin_id": user_id}, {"$set": cascade})
    fresh = await db.users.find_one({"id": user_id}, {"_id": 0, "password_hash": 0})
    return fresh


@api_router.delete("/admin/users/{user_id}")
async def delete_managed_user(user_id: str, user: dict = Depends(get_current_user)):
    role = user.get("role")
    target = await db.users.find_one({"id": user_id})
    if not target:
        raise HTTPException(status_code=404, detail="Pengguna tidak ditemukan")
    if role == "superadmin" and target.get("role") == "admin":
        pass
    elif role == "admin" and target.get("admin_id") == user["id"]:
        pass
    else:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    await db.users.delete_one({"id": user_id})
    return {"success": True}


# ---------------- Document routes ----------------
@api_router.get("/documents/stats")
async def stats(user: dict = Depends(get_current_user)):
    ids = await visible_owner_ids(user)
    match = {} if ids is None else {"owner_id": {"$in": ids}}
    pipeline = [{"$match": match}, {"$group": {"_id": "$type", "count": {"$sum": 1}}}]
    rows = await db.documents.aggregate(pipeline).to_list(100)
    counts = {r["_id"]: r["count"] for r in rows}
    total = sum(counts.values())
    extra = {}
    if user.get("role") in ("admin", "superadmin"):
        extra["managed_users"] = await db.users.count_documents(
            {"role": "admin"} if user["role"] == "superadmin" else {"admin_id": user["id"]}
        )
    return {"total": total, "by_type": counts, **extra}


async def _owner_name_map(owner_ids: List[str]) -> Dict[str, str]:
    users = await db.users.find({"id": {"$in": list(set(owner_ids))}}, {"id": 1, "name": 1, "_id": 0}).to_list(2000)
    return {u["id"]: u.get("name", "") for u in users}


@api_router.get("/documents")
async def list_documents(type: Optional[str] = None, user: dict = Depends(get_current_user)):
    ids = await visible_owner_ids(user)
    q = {} if ids is None else {"owner_id": {"$in": ids}}
    if type:
        q["type"] = type
    docs = await db.documents.find(q, {"_id": 0, "content_html": 0}).sort("updated_at", -1).to_list(1000)
    if user.get("role") in ("admin", "superadmin") and docs:
        names = await _owner_name_map([d["owner_id"] for d in docs])
        for d in docs:
            d["owner_name"] = names.get(d["owner_id"], "")
    return docs


async def _get_visible_doc(doc_id: str, user: dict) -> dict:
    doc = await db.documents.find_one({"id": doc_id}, {"_id": 0})
    if not doc:
        raise HTTPException(status_code=404, detail="Dokumen tidak ditemukan")
    ids = await visible_owner_ids(user)
    if ids is not None and doc["owner_id"] not in ids:
        raise HTTPException(status_code=403, detail="Akses ditolak")
    return doc


@api_router.get("/documents/{doc_id}")
async def get_document(doc_id: str, user: dict = Depends(get_current_user)):
    return await _get_visible_doc(doc_id, user)


@api_router.post("/documents")
async def create_document(data: DocumentInput, user: dict = Depends(get_current_user)):
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()),
        "owner_id": user["id"],
        "type": data.type,
        "title": data.title,
        "meta": data.meta,
        "fields": data.fields,
        "content_html": data.content_html,
        "created_at": now,
        "updated_at": now,
    }
    await db.documents.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.post("/documents/{doc_id}/duplicate")
async def duplicate_document(doc_id: str, user: dict = Depends(get_current_user)):
    src = await _get_visible_doc(doc_id, user)
    now = datetime.now(timezone.utc).isoformat()
    doc = {
        "id": str(uuid.uuid4()),
        "owner_id": user["id"],
        "type": src["type"],
        "title": f"{src['title']} (Salinan)",
        "meta": src.get("meta", {}),
        "fields": src.get("fields", {}),
        "content_html": src.get("content_html"),
        "created_at": now,
        "updated_at": now,
    }
    await db.documents.insert_one(doc)
    doc.pop("_id", None)
    return doc


@api_router.put("/documents/{doc_id}")
async def update_document(doc_id: str, data: DocumentInput, user: dict = Depends(get_current_user)):
    await _get_visible_doc(doc_id, user)
    update = {
        "title": data.title,
        "meta": data.meta,
        "fields": data.fields,
        "content_html": data.content_html,
        "updated_at": datetime.now(timezone.utc).isoformat(),
    }
    await db.documents.update_one({"id": doc_id}, {"$set": update})
    fresh = await db.documents.find_one({"id": doc_id}, {"_id": 0})
    return fresh


@api_router.delete("/documents/{doc_id}")
async def delete_document(doc_id: str, user: dict = Depends(get_current_user)):
    await _get_visible_doc(doc_id, user)
    await db.documents.delete_one({"id": doc_id})
    return {"success": True}


# ---------------- Dropdown options (from data + defaults) ----------------
DEFAULT_OPTIONS = {
    "mataPelajaran": [
        "Bahasa Indonesia", "Matematika", "IPAS", "Pendidikan Pancasila",
        "Bahasa Inggris", "PJOK", "Seni Budaya dan Prakarya (SBdP)",
        "Pendidikan Agama Islam", "Akidah Akhlak", "Fikih", "Al-Qur'an Hadis",
        "Sejarah Kebudayaan Islam", "Bahasa Arab",
    ],
    "kelas": ["1", "2", "3", "4", "5", "6", "7", "8", "9", "10", "11", "12"],
    "fase": ["Fase A", "Fase B", "Fase C", "Fase D", "Fase E", "Fase F"],
    "semester": ["Ganjil", "Genap"],
    "capaianPembelajaran": [],
    "materi": [],
    "alurTujuanPembelajaran": [],
}


@api_router.get("/options")
async def get_options(user: dict = Depends(get_current_user)):
    ids = await visible_owner_ids(user)
    q = {} if ids is None else {"owner_id": {"$in": ids}}
    docs = await db.documents.find(q, {"_id": 0, "fields": 1, "meta": 1}).to_list(1000)
    collected = {k: set() for k in DEFAULT_OPTIONS}
    field_keys = {
        "mataPelajaran": ["mataPelajaran"],
        "kelas": ["kelas"],
        "fase": ["fase"],
        "semester": ["semester"],
        "capaianPembelajaran": ["capaianPembelajaran"],
        "materi": ["materiPokok", "materi"],
        "alurTujuanPembelajaran": ["alurTujuanPembelajaran", "tujuanPembelajaran"],
    }
    for d in docs:
        src = {**(d.get("meta") or {}), **(d.get("fields") or {})}
        for opt, keys in field_keys.items():
            for k in keys:
                v = str(src.get(k, "") or "").strip()
                if v and len(v) < 600:
                    collected[opt].add(v)
    result = {}
    for opt, defaults in DEFAULT_OPTIONS.items():
        vals = sorted(collected[opt])
        merged = list(dict.fromkeys(defaults + vals))
        result[opt] = merged
    return result


# ---------------- DOCX export ----------------
def _add_kop(document: Document, doc: dict):
    f = doc.get("fields") or {}
    m = doc.get("meta") or {}
    sekolah = f.get("namaSekolah") or m.get("namaSekolah") or ""
    alamat = f.get("alamatSekolah") or m.get("alamatSekolah") or ""
    logo = f.get("logoMadrasah") or m.get("logoMadrasah") or ""
    if logo and isinstance(logo, str) and "base64," in logo:
        try:
            b64 = logo.split("base64,", 1)[1]
            img_bytes = base64.b64decode(b64)
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(io.BytesIO(img_bytes), height=Cm(2.0))
        except Exception:
            logger.exception("logo docx render failed")
    if sekolah:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(sekolah.upper())
        r.bold = True
        r.font.size = Pt(14)
    if alamat:
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(alamat)
        r.font.size = Pt(10)
    title = "RENCANA PELAKSANAAN PEMBELAJARAN (RPP)" if doc["type"] == "rpp" else (TYPE_LABEL_BE.get(doc["type"], "DOKUMEN")).upper()
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(12)
    document.add_paragraph("_" * 90)


TYPE_LABEL_BE = {
    "rpp": "RPP", "prota": "Program Tahunan", "prosem": "Program Semester",
    "kktp": "KKTP", "alokasi_waktu": "Alokasi Waktu", "atp": "ATP",
    "asesmen": "Asesmen", "lkpd": "LKPD", "poster": "Poster",
}

RPP_SECTION_LABELS = [
    ("identifikasiPesertaDidik", "Identifikasi Peserta Didik"),
    ("capaianPembelajaran", "Capaian Pembelajaran (CP)"),
    ("dimensiProfilLulusan", "Dimensi Profil Lulusan"),
    ("topikPancaCinta", "Topik Panca Cinta"),
    ("materiIntegrasiKBC", "Materi Integrasi Kurikulum Berbasis Cinta (KBC)"),
    ("pemanfaatanDigital", "Pemanfaatan Digital"),
    ("lintasDisiplin", "Lintas Disiplin Ilmu"),
    ("tujuanPembelajaran", "Tujuan Pembelajaran"),
    ("praktikPedagogik", "Model Pembelajaran"),
    ("metodePembelajaran", "Metode Pembelajaran"),
    ("kemitraan", "Kemitraan Pembelajaran"),
    ("kegiatanAwal", "Langkah Pembelajaran - Kegiatan Awal"),
    ("kegiatanInti", "Langkah Pembelajaran - Kegiatan Inti"),
    ("penutup", "Langkah Pembelajaran - Penutup"),
    ("asesmenAwal", "Asesmen Awal Pembelajaran"),
    ("asesmenProses", "Asesmen Proses Pembelajaran"),
    ("asesmenAkhir", "Asesmen Akhir (Sumatif)"),
    ("rubrikPenilaian", "Rubrik Penilaian"),
]


@api_router.get("/documents/{doc_id}/docx")
async def export_docx(doc_id: str, user: dict = Depends(get_current_user)):
    doc = await _get_visible_doc(doc_id, user)
    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    _add_kop(document, doc)

    if doc["type"] == "rpp":
        f = doc.get("fields") or {}
        identitas = [
            ("Nama Guru", f.get("namaGuru")), ("NIP", f.get("nip")),
            ("Jabatan", f.get("jabatan")), ("Satuan Pendidikan", f.get("namaSekolah")),
            ("Mata Pelajaran", f.get("mataPelajaran")),
            ("Kelas / Fase", " / ".join([x for x in [f.get("kelas"), f.get("fase")] if x])),
            ("Semester", f.get("semester")), ("Materi Pokok", f.get("materiPokok")),
            ("Alokasi Waktu", f.get("alokasiWaktu")), ("Tahun Ajaran", f.get("tahunAjaran")),
        ]
        identitas = [(k, v) for k, v in identitas if v]
        if identitas:
            table = document.add_table(rows=0, cols=2)
            table.style = "Table Grid"
            for k, v in identitas:
                row = table.add_row().cells
                row[0].text = k
                row[1].text = str(v)
            document.add_paragraph("")
        for i, (key, label) in enumerate(RPP_SECTION_LABELS):
            raw = f.get(key)
            if isinstance(raw, list):
                val = ", ".join([str(x).strip() for x in raw if str(x).strip()])
            else:
                val = (raw or "").strip()
            if not val:
                continue
            h = document.add_paragraph()
            r = h.add_run(f"{chr(65 + (i % 26))}. {label}")
            r.bold = True
            document.add_paragraph(val)
        # signature
        document.add_paragraph("")
        sig = document.add_table(rows=1, cols=2)
        left, right = sig.rows[0].cells
        left.text = f"Mengetahui,\nKepala Sekolah\n\n\n\n{f.get('namaKepalaSekolah') or '...........................'}\nNIP. {f.get('nipKepalaSekolah') or '...................'}"
        right.text = f"Guru Mata Pelajaran\n\n\n\n{f.get('namaGuru') or '...........................'}\nNIP. {f.get('nip') or '...................'}"
    else:
        html = doc.get("content_html") or ""
        # strip images (base64) for docx to avoid parser issues
        soup = BeautifulSoup(html, "html.parser")
        for img in soup.find_all("img"):
            img.decompose()
        try:
            HtmlToDocx().add_html_to_document(str(soup), document)
        except Exception:
            document.add_paragraph(soup.get_text("\n"))

    buf = io.BytesIO()
    document.save(buf)
    buf.seek(0)
    safe = re.sub(r"[^A-Za-z0-9_-]+", "_", doc.get("title", "dokumen"))[:60]
    return Response(
        content=buf.read(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{safe}.docx"'},
    )


# ---------------- AI helpers ----------------
RPP_FIELDS = [
    "namaGuru", "nip", "jabatan", "namaSekolah", "alamatSekolah", "mataPelajaran", "kelas",
    "semester", "fase", "materiPokok", "alokasiWaktu", "tahunAjaran",
    "namaKepalaSekolah", "nipKepalaSekolah",
    "identifikasiPesertaDidik", "capaianPembelajaran", "dimensiProfilLulusan",
    "topikPancaCinta", "materiIntegrasiKBC", "pemanfaatanDigital",
    "lintasDisiplin", "tujuanPembelajaran", "praktikPedagogik", "metodePembelajaran", "kemitraan",
    "kegiatanAwal", "kegiatanInti", "penutup",
    "asesmenAwal", "asesmenProses", "asesmenAkhir", "rubrikPenilaian",
]

# Checklist fields (multi-pilih) beserta opsi (jenjang MI, Kurikulum Berbasis Cinta)
CHECKLIST_OPTIONS = {
    "dimensiProfilLulusan": [
        "Keimanan dan Ketakwaan terhadap Tuhan Yang Maha Esa", "Kewargaan",
        "Penalaran Kritis", "Kreativitas", "Kolaborasi", "Kemandirian", "Kesehatan", "Komunikasi",
    ],
    "topikPancaCinta": [
        "Cinta Allah dan Rasul-Nya", "Cinta Ilmu", "Cinta Lingkungan",
        "Cinta Diri dan Sesama", "Cinta Tanah Air",
    ],
    "lintasDisiplin": [
        "Al-Qur'an Hadis", "Akidah Akhlak", "Fikih", "SKI", "Bahasa Arab", "PPKn",
        "Bahasa Indonesia", "Matematika", "IPAS", "Seni Budaya", "PJOK", "Bahasa Inggris", "Muatan Lokal",
    ],
    "praktikPedagogik": [
        "Problem Based Learning", "Project Based Learning", "Discovery Learning",
        "Inquiry Learning", "Cooperative Learning", "Contextual Teaching and Learning",
    ],
    "metodePembelajaran": [
        "Ceramah", "Diskusi", "Tanya Jawab", "Demonstrasi", "Penugasan", "Eksperimen",
        "Simulasi", "Bermain Peran", "Karyawisata", "Drill/Latihan", "Kerja Kelompok",
    ],
    "kemitraan": [
        "Orang Tua/Wali", "Komite Madrasah", "Masyarakat", "Tokoh Agama",
        "Instansi/Lembaga terkait", "Dunia Usaha & Industri", "Alumni",
    ],
}
CHECKLIST_FIELDS = set(CHECKLIST_OPTIONS.keys())


def _clean_json(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(json)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()
    m = re.search(r"\{.*\}", text, re.DOTALL)
    return m.group(0) if m else text


def _clean_html(text: str) -> str:
    text = text.strip()
    text = re.sub(r"^```(html)?", "", text).strip()
    text = re.sub(r"```$", "", text).strip()
    return text


def new_chat(system_message: str) -> LlmChat:
    return LlmChat(
        api_key=EMERGENT_LLM_KEY,
        session_id=str(uuid.uuid4()),
        system_message=system_message,
    ).with_model("gemini", GEMINI_MODEL)


@api_router.post("/ai/extract-rpp")
async def extract_rpp(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    if not EMERGENT_LLM_KEY:
        raise HTTPException(status_code=500, detail="LLM key belum dikonfigurasi")
    filename = (file.filename or "upload").lower()
    ext = filename.rsplit(".", 1)[-1] if "." in filename else ""
    mime_map = {
        "pdf": "application/pdf", "png": "image/png", "jpg": "image/jpeg",
        "jpeg": "image/jpeg", "webp": "image/webp", "heic": "image/heic",
    }
    if ext not in mime_map:
        raise HTTPException(status_code=400, detail="Format tidak didukung. Gunakan PDF atau gambar (PNG/JPG).")
    data = await file.read()
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=f".{ext}")
    tmp.write(data)
    tmp.close()

    system = (
        "Anda asisten ahli kurikulum pendidikan Indonesia (Kurikulum Merdeka & Kurikulum Berbasis Cinta). "
        "Anda membaca dokumen RPP/Modul Ajar yang diunggah lalu mengekstrak isinya menjadi JSON terstruktur."
    )
    keys_desc = ", ".join(RPP_FIELDS)
    checklist_desc = "\n".join(
        [f"  - {k}: pilih dari {opts}" for k, opts in CHECKLIST_OPTIONS.items()]
    )
    prompt = (
        "Baca dokumen terlampir (RPP/Modul Ajar). Ekstrak dan susun kembali isinya menjadi JSON. "
        f"Gunakan PERSIS kunci berikut: {keys_desc}. "
        "Untuk bagian naratif (identifikasi, kegiatan, asesmen, dll) tuliskan isi lengkap yang rapi sebagai STRING (kosongkan '' jika tidak ada). "
        "Untuk kunci berikut kembalikan berupa ARRAY of string, pilih HANYA dari opsi yang tersedia (boleh lebih dari satu, boleh [] jika tidak ada):\n"
        f"{checklist_desc}\n"
        "Jika dokumen berupa format/kerangka kosong, isi dengan konten pembelajaran yang relevan dan lengkap "
        "sesuai mata pelajaran/materi yang terdeteksi. "
        "Balas HANYA dengan objek JSON valid tanpa penjelasan, tanpa ```."
    )
    file_content = FileContentWithMimeType(file_path=tmp.name, mime_type=mime_map[ext])
    try:
        resp = await new_chat(system).send_message(
            UserMessage(text=prompt, file_contents=[file_content])
        )
    except Exception as e:
        logger.exception("extract-rpp failed")
        raise HTTPException(status_code=500, detail=f"Gagal membaca file: {e}")
    finally:
        try:
            os.unlink(tmp.name)
        except Exception:
            pass

    try:
        parsed = json.loads(_clean_json(resp))
    except Exception:
        raise HTTPException(status_code=500, detail="AI tidak mengembalikan data yang valid, coba lagi.")
    fields = {}
    for k in RPP_FIELDS:
        v = parsed.get(k, "")
        if k in CHECKLIST_FIELDS:
            allowed = CHECKLIST_OPTIONS[k]
            if isinstance(v, list):
                items = [str(x).strip() for x in v if str(x).strip()]
            elif isinstance(v, str) and v.strip():
                items = [s.strip() for s in re.split(r"[;,\n]", v) if s.strip()]
            else:
                items = []
            # keep only recognized options (case-insensitive match)
            low = {o.lower(): o for o in allowed}
            fields[k] = [low[i.lower()] for i in items if i.lower() in low]
        else:
            fields[k] = str(v or "")
    # prefill identitas guru dari profil jika kosong
    fields["namaGuru"] = fields["namaGuru"] or user.get("name", "")
    fields["nip"] = fields["nip"] or user.get("nip", "")
    fields["namaSekolah"] = fields["namaSekolah"] or user.get("namaSekolah", "")
    fields["alamatSekolah"] = fields["alamatSekolah"] or user.get("alamatSekolah", "")
    fields["namaKepalaSekolah"] = fields["namaKepalaSekolah"] or user.get("namaKepalaSekolah", "")
    fields["nipKepalaSekolah"] = fields["nipKepalaSekolah"] or user.get("nipKepalaSekolah", "")
    fields["logoMadrasah"] = user.get("logoMadrasah", "")
    return {"fields": fields}


GEN_INSTRUCTIONS = {
    "prota": "Buat PROGRAM TAHUNAN (PROTA). Sajikan tabel berisi: No, Semester, Materi/Lingkup Materi & Capaian Pembelajaran, Alokasi Waktu (JP), Keterangan. Bagi untuk semester Ganjil dan Genap dalam satu tahun ajaran.",
    "prosem": "Buat PROGRAM SEMESTER (PROSEM). Sajikan tabel: No, Tujuan Pembelajaran/Materi, Alokasi Waktu (JP), lalu kolom bulan (Juli–Desember untuk ganjil / Januari–Juni untuk genap) yang dibagi per minggu (1-4) dengan tanda jumlah JP tiap minggu.",
    "kktp": "Buat KRITERIA KETERCAPAIAN TUJUAN PEMBELAJARAN (KKTP). Sajikan tabel: No, Tujuan Pembelajaran, Kriteria Ketercapaian, dan Interval Nilai/Deskripsi (Perlu Bimbingan, Cukup, Baik, Sangat Baik).",
    "alokasi_waktu": "Buat ANALISIS ALOKASI WAKTU. Sajikan tabel perhitungan pekan efektif per bulan dalam satu semester, jumlah pekan tidak efektif, pekan efektif, dan distribusi Jam Pelajaran (JP) untuk tiap materi/Tujuan Pembelajaran.",
    "atp": "Buat ALUR TUJUAN PEMBELAJARAN (ATP). Uraikan dari Capaian Pembelajaran menjadi urutan Tujuan Pembelajaran yang runtut. Sajikan tabel: No, Elemen/CP, Tujuan Pembelajaran, Materi, Alokasi Waktu, Profil Pelajar Pancasila.",
    "asesmen": "Buat INSTRUMEN ASESMEN lengkap: kisi-kisi, 10 soal pilihan ganda (dengan opsi & kunci jawaban), 5 soal essay (dengan kunci/rubrik), serta rubrik penilaian (Sangat Baik/Baik/Cukup/Perlu Bimbingan). Gunakan tabel untuk kisi-kisi dan rubrik.",
    "lkpd": "Buat LEMBAR KERJA PESERTA DIDIK (LKPD). Sertakan: Judul, Identitas (Nama/Kelas), Tujuan Pembelajaran, Petunjuk Pengerjaan, Ringkasan materi singkat, Kegiatan/Langkah kerja, dan soal/pertanyaan latihan untuk siswa. Sediakan garis/kolom isian jawaban.",
    "poster": "Buat KONTEN POSTER MEDIA PEMBELAJARAN yang menarik dan siap cetak A4. Sertakan judul besar, sub-judul, poin-poin kunci materi yang ringkas dan mudah diingat, ajakan/pesan motivasi, serta deskripsi ilustrasi yang disarankan. Gunakan gaya visual (heading besar, daftar berpoin, penekanan).",
}

GEN_LABELS = {
    "prota": "Program Tahunan", "prosem": "Program Semester", "kktp": "KKTP",
    "alokasi_waktu": "Alokasi Waktu", "atp": "ATP", "asesmen": "Asesmen",
    "lkpd": "LKPD", "poster": "Poster",
}


@api_router.post("/ai/generate")
async def generate_document(data: GenerateInput, user: dict = Depends(get_current_user)):
    if not EMERGENT_LLM_KEY:
        raise HTTPException(status_code=500, detail="LLM key belum dikonfigurasi")
    if data.type not in GEN_INSTRUCTIONS:
        raise HTTPException(status_code=400, detail="Jenis dokumen tidak dikenal")

    inp = data.inputs
    ctx_lines = []
    label_map = {
        "mataPelajaran": "Mata Pelajaran", "kelas": "Kelas", "fase": "Fase",
        "semester": "Semester", "tahunAjaran": "Tahun Ajaran",
        "capaianPembelajaran": "Capaian Pembelajaran (CP)",
        "alurTujuanPembelajaran": "Alur Tujuan Pembelajaran (ATP)",
        "materi": "Materi / Lingkup Materi", "alokasiWaktu": "Alokasi Waktu",
        "jumlahMinggu": "Jumlah Minggu Efektif", "catatan": "Catatan Tambahan",
    }
    for k, lbl in label_map.items():
        v = str(inp.get(k, "") or "").strip()
        if v:
            ctx_lines.append(f"- {lbl}: {v}")
    context = "\n".join(ctx_lines) if ctx_lines else "(tidak ada detail tambahan)"

    system = (
        "Anda asisten ahli kurikulum pendidikan Indonesia (Kurikulum Merdeka & Kurikulum Berbasis Cinta) "
        "yang membantu guru menyusun dokumen administrasi pembelajaran. Tulis dalam Bahasa Indonesia yang baku dan lengkap."
    )
    prompt = (
        f"{GEN_INSTRUCTIONS[data.type]}\n\n"
        f"Data acuan dari guru:\n{context}\n\n"
        "Ketentuan output:\n"
        "- Balas HANYA potongan HTML (bukan markdown), TANPA ```html, TANPA tag <html>/<head>/<body>.\n"
        "- Gunakan <h2>, <h3>, <p>, <ul>, <ol>, dan <table> bila perlu.\n"
        "- Untuk tabel gunakan <table><thead><tr><th>...</th></tr></thead><tbody>...</tbody></table> yang rapi.\n"
        "- Konten harus konkret, lengkap, dan langsung bisa dipakai guru.\n"
    )
    try:
        resp = await new_chat(system).send_message(UserMessage(text=prompt))
    except Exception as e:
        logger.exception("generate failed")
        raise HTTPException(status_code=500, detail=f"Gagal membuat dokumen: {e}")

    html = _clean_html(resp)

    if data.type == "poster":
        img_tag = await _generate_poster_image(inp)
        if img_tag:
            html = img_tag + html

    return {"content_html": html}


async def _generate_poster_image(inp: dict) -> Optional[str]:
    """Generate an illustration for the poster using Gemini Nano Banana. Returns an <img> tag with data URL."""
    topic = str(inp.get("materi") or inp.get("mataPelajaran") or "pembelajaran").strip()
    mapel = str(inp.get("mataPelajaran") or "").strip()
    kelas = str(inp.get("kelas") or "").strip()
    prompt = (
        f"Educational poster illustration for elementary/school students about '{topic}'"
        f"{f' ({mapel}, kelas {kelas})' if mapel else ''}. "
        "Colorful, friendly, cartoon flat-illustration style, clean white background, "
        "no text or letters in the image, suitable for a classroom learning poster."
    )
    try:
        chat = LlmChat(
            api_key=EMERGENT_LLM_KEY,
            session_id=str(uuid.uuid4()),
            system_message="You are an assistant that generates educational illustrations.",
        ).with_model("gemini", GEMINI_IMAGE_MODEL).with_params(modalities=["image", "text"])
        _text, images = await chat.send_message_multimodal_response(UserMessage(text=prompt))
        if images:
            img = images[0]
            mime = img.get("mime_type", "image/png")
            data = img.get("data", "")
            if data:
                return (
                    f'<div style="text-align:center;margin:0 0 16px;">'
                    f'<img src="data:{mime};base64,{data}" '
                    f'style="max-width:70%;height:auto;border-radius:8px;" alt="Ilustrasi poster" />'
                    f'</div>'
                )
    except Exception:
        logger.exception("poster image generation failed")
    return None


# ---------------- App wiring ----------------
@api_router.get("/")
async def root():
    return {"message": "Generator Pembelajaran API"}


app.include_router(api_router)

app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=os.environ.get('CORS_ORIGINS', '*').split(','),
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("id")
    await db.users.create_index("admin_id")
    await db.documents.create_index("owner_id")

    # seed super admin
    sa = await db.users.find_one({"email": SUPERADMIN_EMAIL.lower()})
    if not sa:
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "name": "Super Admin",
            "email": SUPERADMIN_EMAIL.lower(),
            "password_hash": hash_password(SUPERADMIN_PASSWORD),
            "role": "superadmin",
            "admin_id": None,
            "created_by": None,
            "nip": "", "jabatan": "Super Admin",
            "namaSekolah": "", "alamatSekolah": "",
            "namaKepalaSekolah": "", "nipKepalaSekolah": "", "logoMadrasah": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    elif not verify_password(SUPERADMIN_PASSWORD, sa["password_hash"]):
        await db.users.update_one({"email": SUPERADMIN_EMAIL.lower()}, {"$set": {
            "password_hash": hash_password(SUPERADMIN_PASSWORD), "role": "superadmin"}})

    # seed demo guru
    demo_email = "guru@demo.com"
    if not await db.users.find_one({"email": demo_email}):
        await db.users.insert_one({
            "id": str(uuid.uuid4()),
            "name": "Ibu Eka Novita Sari",
            "email": demo_email,
            "password_hash": hash_password("guru123"),
            "role": "guru",
            "admin_id": None,
            "created_by": None,
            "nip": "199001012020122001",
            "jabatan": "Guru Kelas",
            "namaSekolah": "MI Miftahul Jannah",
            "alamatSekolah": "Jl. Pendidikan No. 1, Kec. Sukamaju, Kab. Bogor, Jawa Barat",
            "namaKepalaSekolah": "H. Zainur Ridho, S.Pd.I",
            "nipKepalaSekolah": "197505052005011003",
            "logoMadrasah": "",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
